"""Automated review engine.

Given a template (its field spec + original file) and a candidate's submitted
file, this runs the automated checks and produces:

* extracted ``field_values``,
* a :class:`ReviewReport` (per-check pass/fail with human messages),
* a :class:`DiffSummary` (submitted vs original, per field),
* an overall :class:`ReviewStatus` and a simple score.

Status rules:
* any error-severity failure  -> ``failed``
* otherwise any warning        -> ``needs_review``
* otherwise                    -> ``passed``
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl.utils import get_column_letter

from app.models.enums import FileType, ReviewStatus
from app.schemas.field_spec import FieldDefinition, FieldSpec, FieldValueType
from app.schemas.submission import CheckResult, DiffSummary, FieldDiff, ReviewReport
from app.services import docx_service, xlsx_service


@dataclass
class ReviewOutcome:
    values: dict[str, object]
    report: ReviewReport
    diff: DiffSummary
    score: float | None


def _json_safe(value: object):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _is_empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _as_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _strip_html(value: object) -> str:
    return re.sub(r"<[^>]+>", "", str(value)).strip()


def _validate_field(fd: FieldDefinition, value: object) -> list[CheckResult]:
    checks: list[CheckResult] = []
    label = fd.label

    # Required.
    if _is_empty(value):
        if fd.required:
            checks.append(
                CheckResult(
                    id=f"field:{fd.id}:required",
                    label=f"{label} — required",
                    passed=False,
                    severity="error",
                    message=f"'{label}' is required but was left empty.",
                )
            )
        return checks  # nothing else to validate on an empty value

    checks.append(
        CheckResult(
            id=f"field:{fd.id}:present",
            label=f"{label} — provided",
            passed=True,
            severity="info",
            message=f"'{label}' was filled in.",
        )
    )

    # Numeric range/type.
    if fd.type in (FieldValueType.number, FieldValueType.integer):
        num = _as_number(value)
        if num is None:
            checks.append(
                CheckResult(
                    id=f"field:{fd.id}:type",
                    label=f"{label} — numeric",
                    passed=False,
                    severity="error",
                    message=f"'{label}' must be a number, got '{value}'.",
                )
            )
        else:
            if fd.type is FieldValueType.integer and not float(num).is_integer():
                checks.append(
                    CheckResult(
                        id=f"field:{fd.id}:integer",
                        label=f"{label} — whole number",
                        passed=False,
                        severity="error",
                        message=f"'{label}' must be a whole number, got {num}.",
                    )
                )
            if fd.min is not None and num < fd.min:
                checks.append(
                    CheckResult(
                        id=f"field:{fd.id}:min",
                        label=f"{label} — minimum",
                        passed=False,
                        severity="error",
                        message=f"'{label}' must be ≥ {fd.min}, got {num}.",
                    )
                )
            if fd.max is not None and num > fd.max:
                checks.append(
                    CheckResult(
                        id=f"field:{fd.id}:max",
                        label=f"{label} — maximum",
                        passed=False,
                        severity="error",
                        message=f"'{label}' must be ≤ {fd.max}, got {num}.",
                    )
                )

    # String constraints (evaluate on visible text for rich fields).
    text = _strip_html(value) if fd.rich_text else str(value)
    if fd.min_length is not None and len(text) < fd.min_length:
        checks.append(
            CheckResult(
                id=f"field:{fd.id}:min_length",
                label=f"{label} — minimum length",
                passed=False,
                severity="error",
                message=f"'{label}' must be at least {fd.min_length} characters.",
            )
        )
    if fd.max_length is not None and len(text) > fd.max_length:
        checks.append(
            CheckResult(
                id=f"field:{fd.id}:max_length",
                label=f"{label} — maximum length",
                passed=False,
                severity="error",
                message=f"'{label}' must be at most {fd.max_length} characters.",
            )
        )
    if fd.pattern and not re.search(fd.pattern, text):
        checks.append(
            CheckResult(
                id=f"field:{fd.id}:pattern",
                label=f"{label} — format",
                passed=False,
                severity="error",
                message=f"'{label}' does not match the required format.",
            )
        )
    if fd.allowed_values and str(value) not in fd.allowed_values:
        checks.append(
            CheckResult(
                id=f"field:{fd.id}:allowed",
                label=f"{label} — allowed value",
                passed=False,
                severity="error",
                message=f"'{label}' must be one of: {', '.join(fd.allowed_values)}.",
            )
        )
    return checks


def _extract(file_type: FileType, spec: FieldSpec, data: bytes):
    """Return (values, uncaptured_ids)."""
    if file_type is FileType.docx:
        res = docx_service.extract_field_values(data, spec.for_word())
        return res.values, set(res.uncaptured)
    values = xlsx_service.extract_field_values(data, spec.for_excel())
    return values, set()


def review_submission(
    file_type: FileType,
    field_spec: FieldSpec,
    submitted_bytes: bytes,
    original_bytes: bytes,
    provided_values: dict[str, object] | None = None,
) -> ReviewOutcome:
    """Review a submission.

    ``provided_values`` is supplied for in-browser edits (the authoritative values
    the candidate entered). When given, we validate those directly instead of
    re-extracting them from the generated file. For uploads it is ``None`` and we
    extract values from the submitted file.
    """
    fields = field_spec.for_word() if file_type is FileType.docx else field_spec.for_excel()

    if provided_values is not None:
        submitted_values: dict[str, object] = {
            fd.id: provided_values.get(fd.id) for fd in fields
        }
        uncaptured: set[str] = set()
    else:
        submitted_values, uncaptured = _extract(file_type, field_spec, submitted_bytes)

    try:
        original_values, _ = _extract(file_type, field_spec, original_bytes)
    except Exception:  # noqa: BLE001 - original diff is best-effort
        original_values = {}

    checks: list[CheckResult] = []
    fields_passed = 0

    for fd in fields:
        value = submitted_values.get(fd.id)
        if fd.id in uncaptured:
            checks.append(
                CheckResult(
                    id=f"field:{fd.id}:manual",
                    label=f"{fd.label} — manual review",
                    passed=True,
                    severity="warning",
                    message=(
                        f"'{fd.label}' was edited but its value could not be captured "
                        "automatically; please review manually."
                    ),
                )
            )
            continue
        field_checks = _validate_field(fd, value)
        checks.extend(field_checks)
        if not any(c.severity == "error" and not c.passed for c in field_checks):
            fields_passed += 1

    # File integrity: for uploaded Excel, flag changes to cells outside the
    # editable set. Skipped for in-browser edits (we generated the file).
    if file_type is FileType.xlsx and provided_values is None:
        checks.extend(_detect_unexpected_xlsx_changes(field_spec, original_bytes, submitted_bytes))

    # Build the diff.
    diffs: list[FieldDiff] = []
    changed = 0
    for fd in fields:
        sub = _json_safe(submitted_values.get(fd.id))
        orig = _json_safe(original_values.get(fd.id))
        is_changed = sub != orig
        changed += 1 if is_changed else 0
        diffs.append(
            FieldDiff(field_id=fd.id, label=fd.label, original=orig, submitted=sub, changed=is_changed)
        )
    diff = DiffSummary(fields=diffs, changed_count=changed)

    failed = [c for c in checks if not c.passed and c.severity == "error"]
    passed = [c for c in checks if c.passed and c.severity != "warning"]

    if failed:
        status = ReviewStatus.failed
    elif any(c.severity == "warning" for c in checks):
        status = ReviewStatus.needs_review
    else:
        status = ReviewStatus.passed

    summary = _build_summary(status, len(fields), fields_passed, len(failed))
    report = ReviewReport(
        status=status,
        summary=summary,
        checks=checks,
        passed_count=len(passed),
        failed_count=len(failed),
        warning_count=sum(1 for c in checks if c.severity == "warning"),
    )
    score = round(100 * fields_passed / len(fields), 1) if fields else None
    return ReviewOutcome(
        values={k: _json_safe(v) for k, v in submitted_values.items()},
        report=report,
        diff=diff,
        score=score,
    )


def _detect_unexpected_xlsx_changes(
    spec: FieldSpec, original_bytes: bytes, submitted_bytes: bytes
) -> list[CheckResult]:
    editable: set[tuple[str, str]] = {
        (fd.sheet, fd.cell.upper()) for fd in spec.for_excel() if fd.sheet and fd.cell
    }
    try:
        wb_o = xlsx_service._open(original_bytes, data_only=False)
        wb_s = xlsx_service._open(submitted_bytes, data_only=False)
    except Exception:  # noqa: BLE001
        return []
    try:
        unexpected: list[str] = []
        for name in wb_o.sheetnames:
            if name not in wb_s.sheetnames:
                unexpected.append(f"sheet '{name}' removed")
                continue
            ws_o, ws_s = wb_o[name], wb_s[name]
            max_row = min(max(ws_o.max_row or 1, ws_s.max_row or 1), 500)
            max_col = min(max(ws_o.max_column or 1, ws_s.max_column or 1), 100)
            for r in range(1, max_row + 1):
                for c in range(1, max_col + 1):
                    addr = f"{get_column_letter(c)}{r}"
                    if (name, addr) in editable:
                        continue
                    if ws_o.cell(r, c).value != ws_s.cell(r, c).value:
                        unexpected.append(f"{name}!{addr}")
        if unexpected:
            preview = ", ".join(unexpected[:8]) + (" …" if len(unexpected) > 8 else "")
            return [
                CheckResult(
                    id="file:unexpected_changes",
                    label="Unexpected changes",
                    passed=False,
                    severity="warning",
                    message=f"Cells outside the editable fields were changed: {preview}",
                )
            ]
        return []
    finally:
        wb_o.close()
        wb_s.close()


def _build_summary(status: ReviewStatus, total: int, passed: int, failures: int) -> str:
    if status is ReviewStatus.passed:
        return f"All checks passed ({passed}/{total} fields valid)."
    if status is ReviewStatus.failed:
        return f"{failures} check(s) failed. {passed}/{total} fields valid."
    return f"Submission needs manual review. {passed}/{total} fields auto-validated."
