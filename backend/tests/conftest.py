"""Shared test fixtures and in-memory file builders.

We build sample .docx/.xlsx files programmatically so the repo carries no binary
fixtures and every test is self-contained and deterministic.
"""
from __future__ import annotations

import os

# Configure an isolated test environment BEFORE any app module is imported
# (settings are cached on first access).
os.environ.setdefault("DATABASE_URL", "sqlite:///./data/test.db")
os.environ.setdefault("STORAGE_BACKEND", "local")
os.environ.setdefault("LOCAL_STORAGE_DIR", "./data/test_files")
os.environ.setdefault("ADMIN_EMAIL", "admin@example.com")
os.environ.setdefault("ADMIN_PASSWORD", "test-pass-123")
os.environ.setdefault("JWT_SECRET", "test-secret")
os.environ.setdefault("ENVIRONMENT", "test")
# Tests assume auth is enforced (override any local .env that disables it).
os.environ["AUTH_ENABLED"] = "true"

import io
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook


# --------------------------------------------------------------------------- #
# Word builders
# --------------------------------------------------------------------------- #
def build_docx(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def add_content_control(doc: Document, tag: str, text: str) -> None:
    """Append a block-level content control (SDT) with a tag and inner text."""
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    tag_el = OxmlElement("w:tag")
    tag_el.set(qn("w:val"), tag)
    sdt_pr.append(tag_el)
    sdt.append(sdt_pr)
    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    sdt_content.append(run)
    sdt.append(sdt_content)
    doc.element.body.append(sdt)


def build_docx_with_content_control(tag: str, text: str) -> bytes:
    doc = Document()
    doc.add_paragraph("Assessment")
    add_content_control(doc, tag, text)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def build_macro_docx() -> bytes:
    """A structurally valid .docx that also embeds a vbaProject.bin (macro)."""
    base = build_docx(["hello"])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.namelist():
            dst.writestr(item, src.read(item))
        dst.writestr("word/vbaProject.bin", b"\x00\x01macrobytes")
    return out.getvalue()


# --------------------------------------------------------------------------- #
# Excel builders
# --------------------------------------------------------------------------- #
def build_xlsx(cells: dict[str, object], sheet_name: str = "Sheet1") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for addr, value in cells.items():
        ws[addr] = value
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# --------------------------------------------------------------------------- #
# Fixtures
# --------------------------------------------------------------------------- #
@pytest.fixture
def sample_docx() -> bytes:
    return build_docx(["Dear {{full_name}},", "Please write your motivation: {{motivation}}"])


@pytest.fixture
def sample_xlsx() -> bytes:
    return build_xlsx({"A1": "Name", "B1": "", "A2": "Score", "B2": 0})
